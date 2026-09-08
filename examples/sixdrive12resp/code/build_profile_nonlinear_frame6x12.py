# -*- coding: utf-8 -*-
"""
build_profile_nonlinear_frame6x12.py

Derives sdynpy_frame6x12_profile_nonlinear.xlsx from the already-generated
sdynpy_frame6x12_profile.xlsx (built by build_sdynpy_demo_frame6x12.py,
which needs the `sdynpy` conda env). This script itself only needs
openpyxl -- no sdynpy -- since all it does is copy the linear profile and
repoint the Hardware sheet at the nonlinear system file.

Everything else (channel table, control law + parameters, control
channels, spec file, control-phase averaging, etc.) is identical between
the linear and nonlinear runs -- same physical rig, same 8 control
channels (nodes 7-14), same match_trace_pseudoinverse_pi tuning. Only the
hardware being driven differs: SDynPy System Integration (index 6) with
the linear system file, vs. SDynPy Nonlinear System Integration (index 7)
with the nonlinear all-modes system file.

Run this AFTER sdynpy_frame6x12_profile.xlsx exists (i.e. after running
build_sdynpy_demo_frame6x12.py at least once with the 2026-09-04 settings).

    python build_profile_nonlinear_frame6x12.py

Output: ../results/sdynpy_frame6x12_profile_nonlinear.xlsx
"""

import os
import shutil
import openpyxl as opxl

# This script's own __file__ location is not a reliable basis for the path
# written INTO the profile: this script is meant to be runnable either on
# the machine that actually runs Rattlesnake, or (as happened 2026-09-04)
# through a filesystem bridge that mounts the same folder under a different
# local path -- __file__ would then resolve to the *bridge's* path, and a
# profile carrying that path would fail to find the hardware file when
# Rattlesnake itself later loads it on the real machine. Instead, derive
# the nonlinear system file's path from the LINEAR profile's own Hardware
# File value, which is already correct for wherever Rattlesnake actually
# runs (it was written by build_sdynpy_demo_frame6x12.py running natively
# there) -- swap only the filename, keep that directory.
RESULTS_DIR = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "results"))
LINEAR_PROFILE = os.path.join(RESULTS_DIR, "sdynpy_frame6x12_profile.xlsx")
NONLINEAR_PROFILE = os.path.join(RESULTS_DIR, "sdynpy_frame6x12_profile_nonlinear.xlsx")
NONLINEAR_SYSTEM_FILENAME = "sdynpy_frame6x12_system_nonlinear_allmodes.npz"
# Local (this-filesystem-view) path, used only for the existence check below.
NONLINEAR_SYSTEM_FILE_LOCAL = os.path.join(RESULTS_DIR, NONLINEAR_SYSTEM_FILENAME)

NONLINEAR_HARDWARE_INDEX = 7  # SDynPy Nonlinear System Integration (index 6 = linear)

if not os.path.exists(LINEAR_PROFILE):
    raise FileNotFoundError(
        f"{LINEAR_PROFILE} does not exist yet -- run build_sdynpy_demo_frame6x12.py "
        "first (needs the sdynpy conda env) to generate the linear profile this "
        "script derives the nonlinear one from."
    )
if not os.path.exists(NONLINEAR_SYSTEM_FILE_LOCAL):
    raise FileNotFoundError(
        f"{NONLINEAR_SYSTEM_FILE_LOCAL} does not exist -- run "
        "build_nonlinear_frf_system_allmodes.py first to build it."
    )

shutil.copyfile(LINEAR_PROFILE, NONLINEAR_PROFILE)

workbook = opxl.load_workbook(NONLINEAR_PROFILE)
hardware_ws = workbook['Hardware']

old_index = hardware_ws.cell(1, 2).value
old_file = hardware_ws.cell(2, 2).value

# Derive the path to write from the linear profile's own (already-correct,
# machine-native) directory rather than from this script's __file__.
linear_dir = old_file.rsplit('/', 1)[0] if '/' in old_file else old_file.rsplit('\\', 1)[0]
sep = '/' if '/' in old_file else '\\'
nonlinear_system_file = linear_dir + sep + NONLINEAR_SYSTEM_FILENAME

hardware_ws.cell(1, 2, NONLINEAR_HARDWARE_INDEX)      # Hardware Type
hardware_ws.cell(2, 2, nonlinear_system_file)         # Hardware File

workbook.save(NONLINEAR_PROFILE)

print(f"Wrote {NONLINEAR_PROFILE}")
print(f"  Hardware Type: {old_index} -> {NONLINEAR_HARDWARE_INDEX} (SDynPy Nonlinear System Integration)")
print(f"  Hardware File: {old_file} -> {nonlinear_system_file}")
print("Everything else (channel table, control law, control channels, spec, "
      "averaging) copied unchanged from the linear profile.")
